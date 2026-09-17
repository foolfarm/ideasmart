import csv
import json
import os
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

if len(sys.argv) != 4:
    raise SystemExit("Uso: python3 create-newsletter-subscribers-xlsx.py <input.csv> <summary.json> <output.xlsx>")

csv_path, summary_path, output_path = sys.argv[1:]

with open(summary_path, "r", encoding="utf-8") as summary_file:
    summary = json.load(summary_file)

with open(csv_path, newline="", encoding="utf-8-sig") as csv_file:
    reader = csv.DictReader(csv_file)
    subscribers = [row for row in reader if row.get("Email")]

workbook = Workbook()
worksheet = workbook.active
worksheet.title = "Email iscritti"
worksheet.sheet_view.showGridLines = False

primary = "151515"
accent = "E04B35"
light = "F5F2EE"
gray = "666666"
white = "FFFFFF"
border_side = Side(style="thin", color="D9D4CF")

worksheet.column_dimensions["A"].width = 3
worksheet.column_dimensions["B"].width = 46
worksheet.merge_cells("B2:E2")
worksheet["B2"] = "PROOFPRESS — ELENCO ISCRITTI NEWSLETTER"
worksheet["B2"].font = Font(name="Georgia", size=16, bold=True, color=primary)
worksheet["B2"].alignment = Alignment(horizontal="left", vertical="center")
worksheet.row_dimensions[2].height = 30

worksheet.merge_cells("B3:E3")
worksheet["B3"] = f"{len(subscribers):,}".replace(",", ".") + " iscritti attivi · esportazione " + datetime.fromisoformat(summary["generatedAt"].replace("Z", "+00:00")).astimezone().strftime("%d/%m/%Y %H:%M")
worksheet["B3"].font = Font(name="Calibri", size=10, italic=True, color=gray)
worksheet["B3"].alignment = Alignment(horizontal="left", vertical="center")
worksheet.row_dimensions[3].height = 18

worksheet.merge_cells("B4:E4")
worksheet["B4"] = "Nota: il database degli iscritti non conserva numeri di telefono. La colonna è presente per future integrazioni e risulta vuota in questo export."
worksheet["B4"].font = Font(name="Calibri", size=9, italic=True, color=gray)
worksheet["B4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
worksheet.row_dimensions[4].height = 30

header_row = 6
headers = ["Nome", "Cognome", "Telefono", "Email"]
for column, value in enumerate(headers, start=2):
    header = worksheet.cell(row=header_row, column=column, value=value)
    header.font = Font(name="Georgia", size=11, bold=True, color=white)
    header.fill = PatternFill("solid", fgColor=primary)
    header.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    header.border = Border(top=border_side, bottom=Side(style="medium", color=primary), left=border_side if column == 2 else Side(style=None), right=border_side if column == 5 else Side(style=None))
worksheet.row_dimensions[header_row].height = 24

for row_number, subscriber in enumerate(subscribers, start=header_row + 1):
    values = [subscriber.get("Nome", ""), subscriber.get("Cognome", ""), subscriber.get("Telefono", ""), subscriber.get("Email", "")]
    for column, value in enumerate(values, start=2):
        cell = worksheet.cell(row=row_number, column=column, value=value)
        cell.font = Font(name="Calibri", size=10, color=primary)
        cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        cell.fill = PatternFill("solid", fgColor="FFFFFF" if row_number % 2 == 0 else light)
        cell.border = Border(left=border_side if column == 2 else Side(style=None), right=border_side if column == 5 else Side(style=None), bottom=border_side if row_number == header_row + len(subscribers) else Side(style="hair", color="E8E4DF"))
    worksheet.row_dimensions[row_number].height = 18

worksheet.column_dimensions["B"].width = 22
worksheet.column_dimensions["C"].width = 28
worksheet.column_dimensions["D"].width = 20
worksheet.column_dimensions["E"].width = 46
worksheet.freeze_panes = "B7"
worksheet.auto_filter.ref = f"B{header_row}:E{header_row + len(subscribers)}"
worksheet.sheet_properties.pageSetUpPr.fitToPage = True
worksheet.page_setup.fitToWidth = 1
worksheet.page_setup.fitToHeight = 0
worksheet.page_margins.left = 0.25
worksheet.page_margins.right = 0.25
worksheet.page_margins.top = 0.5
worksheet.page_margins.bottom = 0.5
worksheet.sheet_view.zoomScale = 95

os.makedirs(os.path.dirname(output_path), exist_ok=True)
workbook.save(output_path)
print(json.dumps({"outputXlsx": output_path, "records": len(subscribers)}, ensure_ascii=False))
